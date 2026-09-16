"""시험 결과를 엑셀(.xlsx) 파일 하나에 계속 누적 저장한다.

DB는 나중에(기능 검증 후) 별도로 다시 붙이기로 하고, 그때까지는 배포 대상 PC에 DB 설치
없이도 결과를 남길 수 있게 하기 위한 조치(사용자 요청, 2026-09-16). 처음엔 세션마다 별도
텍스트 파일(write_session_txt)로 저장했었는데, 여러 세션을 한 눈에 비교/정렬하려면 파일이
따로따로 흩어져 있는 것보다 엑셀 한 파일에 계속 쌓이는 편이 훨씬 낫다는 요청에 따라 이걸로
대체한다.

세션(부품 ID 하나)당 6행(시작점 + 이동량/데드클릭/드리프트/쉬프트/백래쉬) x (날짜/부품ID
공통 + 상/하/좌/우 + 결과 + 종합 결과) 블록을 기존 파일 맨 아래에 이어붙인다. 파일이 없으면
헤더와 함께 새로 만든다. `InspectionRepository`(core/data/repository.py)는 그대로 남겨둬서
나중에 DB를 다시 붙일 때 같은 InspectionSession을 그대로 쓸 수 있다.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from core.inspection.models import CheckType, DirectionTestResult, InspectionSession, TravelDirection, Verdict

_DIRECTION_ORDER = (TravelDirection.UP, TravelDirection.DOWN, TravelDirection.LEFT, TravelDirection.RIGHT)
# 시작점은 CheckType이 아니라 DirectionTestResult.start_point_moa를 보여주는 가상의 행 -
# 종합 결과표(app/views/stage2_travel_test_view.py)와 동일한 패턴.
_START_POINT_ROW = "start_point"
_ROW_LABELS = {
    _START_POINT_ROW: "시작점",
    CheckType.TRAVEL_AMOUNT: "이동량",
    CheckType.DEAD_CLICK: "데드클릭",
    CheckType.DRIFT: "드리프트",
    CheckType.SHIFT: "쉬프트",
    CheckType.BACKLASH: "백래쉬",
}
_ROW_ORDER = (
    _START_POINT_ROW,
    CheckType.TRAVEL_AMOUNT,
    CheckType.DEAD_CLICK,
    CheckType.DRIFT,
    CheckType.SHIFT,
    CheckType.BACKLASH,
)
_HEADER = ["날짜/시간", "부품 ID", "항목", "상", "하", "좌", "우", "결과", "종합 결과"]
_MERGED_COLUMNS = (1, 2, 9)  # 날짜/부품 ID/종합 결과 - 세션당 한 번만 보이면 되는 열

DEFAULT_XLSX_PATH = Path(__file__).resolve().parents[2] / "reports" / "results.xlsx"


def append_session_xlsx(session: InspectionSession, path: Path | None = None) -> Path:
    """세션 결과를 엑셀 파일 맨 아래에 새 블록으로 추가하고, 저장된 파일 경로를 반환한다."""
    target = Path(path) if path is not None else DEFAULT_XLSX_PATH
    target.parent.mkdir(parents=True, exist_ok=True)

    if target.exists():
        wb = load_workbook(target)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "결과"
        ws.append(_HEADER)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col_idx in range(1, len(_HEADER) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 14
        ws.column_dimensions[get_column_letter(1)].width = 19  # "날짜/시간"은 더 길다(YYYY-MM-DD HH:MM:SS)

    results_by_direction = {r.direction: r for r in session.direction_results}
    # 날짜만으로는 같은 날 여러 번 시험한 세션을 구분할 수 없다는 지적(2026-09-16)에 따라
    # 시:분:초까지 같이 기록한다.
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    start_row = ws.max_row + 1

    for row_kind in _ROW_ORDER:
        row: list[object] = [timestamp, session.scope_id, _ROW_LABELS[row_kind]]
        row_statuses: list[str] = []
        for direction in _DIRECTION_ORDER:
            result = results_by_direction.get(direction)
            if row_kind == _START_POINT_ROW:
                row.append(_start_point_text(result))
                continue
            text, status = _cell_for(result, row_kind)
            row.append(text)
            if status is not None:
                row_statuses.append(status)

        if row_kind == _START_POINT_ROW or not row_statuses:
            row_result = "-"
        elif "불량" in row_statuses:
            row_result = "불량"
        else:
            row_result = "합격"
        row.append(row_result)
        row.append("")  # 종합 결과 - 아래에서 이 세션의 첫 행에만 채움
        ws.append(row)

    end_row = ws.max_row
    for col in _MERGED_COLUMNS:
        ws.merge_cells(start_row=start_row, start_column=col, end_row=end_row, end_column=col)
        ws.cell(row=start_row, column=col).alignment = Alignment(vertical="center", horizontal="center")
    ws.cell(row=start_row, column=9, value=session.overall_verdict.value)

    wb.save(target)
    return target


def _start_point_text(result: DirectionTestResult | None) -> str:
    if result is None or result.start_point_moa is None:
        return "-"
    x, y = result.start_point_moa
    return f"({x:.1f}, {y:.1f})"


def _cell_for(result: DirectionTestResult | None, check_type: CheckType) -> tuple[str, str | None]:
    """결과표 한 칸(방향 x 항목)의 표시 문자열과 판정("합격"/"불량"/None=미평가)."""
    if result is None:
        return "-", None

    check = next((c for c in result.check_results if c.check_type == check_type), None)

    if check_type == CheckType.DEAD_CLICK:
        # DEAD_CLICK은 데드클릭으로 표시(flagged=True)했을 때만 CheckResult가 존재한다
        # (TravelTestStateMachine.set_dead_click 참고) - 없으면 "X"(정상), 있으면 "O"(불량).
        return ("O", Verdict.FAIL.value) if check is not None else ("X", None)

    if check is None:
        return "-", None  # 이동량/쉬프트/드리프트 불합격으로 원점 복귀 없이 종료돼 백래쉬만 없는 경우 등
    if check.measured_value is None:
        value_text = "-"
    elif check.point_moa is not None:
        px, py = check.point_moa
        value_text = f"{check.measured_value:.1f}({px:.1f}, {py:.1f})"
    else:
        value_text = f"{check.measured_value:.1f}"
    return value_text, check.status.value
